"""Format-specific loaders. Each takes a file path and returns (text, metadata)."""

from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader


def _metadata(path: Path) -> dict:
    return {"source": path.name, "file_type": path.suffix.lstrip(".").lower()}


def load_pdf(path: Path) -> tuple[str, dict]:
    reader = PdfReader(str(path))
    text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    return text, _metadata(path)


def load_html(path: Path) -> tuple[str, dict]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    text = soup.get_text(separator="\n\n")
    return text, _metadata(path)


def load_docx(path: Path) -> tuple[str, dict]:
    doc = DocxDocument(str(path))
    text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return text, _metadata(path)


def load_pptx(path: Path) -> tuple[str, dict]:
    prs = Presentation(str(path))
    slides_text = []
    for slide in prs.slides:
        slide_text = "\n".join(
            shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
        )
        if slide_text.strip():
            slides_text.append(slide_text)
    return "\n\n".join(slides_text), _metadata(path)


def load_txt(path: Path) -> tuple[str, dict]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return text, _metadata(path)
