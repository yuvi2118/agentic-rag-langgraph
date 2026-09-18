from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from bs4 import BeautifulSoup

def process_document(path: str) -> list[dict]:
    file_path = Path(path)
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        text = extract_pdf(file_path)
    elif extension == ".docx":
        text = extract_docx(file_path)
    elif extension == ".pptx":
        text = extract_pptx(file_path)
    elif extension == ".html" or extension == ".htm":
        text = extract_html(file_path)
    elif extension == ".txt":
        text = file_path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {extension}")
    chunks = chunk_text(text , source=file_path.name)
    return chunks

def extract_pdf(file_path : Path) -> str:
    reader = PdfReader(str(file_path))
    text = ""
    for page in reader.pages:
        text += page.extract_text() or ""  # extract_text() can return None for an image-only page
        text += "\n\n"
    return text

def extract_docx(file_path : Path) -> str:
    doc = Document(str(file_path))
    text = ""
    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text + "\n\n"
    return text


def extract_pptx(file_path : Path) ->str:
    prs = Presentation(str(file_path))
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:  # not every shape has text (images, lines, etc.)
                text += shape.text_frame.text + "\n"
        text += "\n"
    return text


def extract_html(file_path : Path) ->str:
    html_content = file_path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(html_content,"html.parser")
    return soup.get_text(separator="\n\n")  # strips tags, keeps paragraph breaks so chunk_text can split on them

def chunk_text(text : str, source:str , max_chars : int = 1500) -> list[dict]:
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

    chunks = []
    current_chunk = ""

    for para in paragraphs:
        # keep packing paragraphs into the current chunk until it would exceed max_chars
        if len(current_chunk) + len(para) <= max_chars:
            current_chunk += para + "\n\n"
        else:
            if current_chunk:
                chunks.append({"text": current_chunk.strip(), "source": source})
            current_chunk = para + "\n\n"  # start the next chunk with the paragraph that didn't fit

    if current_chunk:  # last chunk never hits the "else" branch above, so add it here
        chunks.append({"text": current_chunk.strip(), "source": source})

    return chunks


MAX_FILE_SIZE_MB = 20


def ingest_folder(folder_path: str) -> list[dict]:
    all_chunks = []
    success_count = 0
    failed_files = []

    for file_path in Path(folder_path).iterdir():
        if not file_path.is_file():  # skip subfolders
            continue

        size_mb = file_path.stat().st_size / (1024 * 1024)
        if size_mb > MAX_FILE_SIZE_MB:
            print(f"Skipping {file_path.name}: too large ({size_mb:.1f} MB)")
            failed_files.append(file_path.name)
            continue

        try:
            chunks = process_document(str(file_path))
            all_chunks.extend(chunks)
            success_count += 1
            print(f"OK: {file_path.name} -> {len(chunks)} chunks")
        except Exception as e:  # one bad file shouldn't stop the whole batch
            print(f"FAILED: {file_path.name} -> {e}")
            failed_files.append(file_path.name)

    print(f"\nDone. {success_count} succeeded, {len(failed_files)} failed.")
    if failed_files:
        print("Failed files:", failed_files)

    return all_chunks
